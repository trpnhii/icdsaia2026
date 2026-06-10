"""
Generate paper-ready experimental results artifacts.

Artifacts written to:
  output_data/experimental_results/
    experimental_setup.json
    experimental_setup.md
    table1_anomaly_baseline_comparison.csv
    table1b_event_ranking_summary.csv
    table1_primary_metrics.csv
    table1c_event_case_studies.csv
    table2_inventory_cost_benefit.csv
    table3_financial_impact.csv
    event_ranking_diagnostics.csv
    figure1_stl_decomposition.png
    figure2_baseline_comparison.png
    results_summary.md

This script is designed to sit on top of the existing pipeline outputs and
normalize minor schema differences between older and newer preprocessing code.
"""

from __future__ import annotations

import argparse
import json
import math
import re
from dataclasses import dataclass
from pathlib import Path

import matplotlib.dates as mdates
import matplotlib.pyplot as plt
import numpy as np
import pandas as pd
from statsmodels.tsa.seasonal import STL


PROJECT_ROOT = Path(__file__).resolve().parents[1]

TRANSFORMED_PATH = PROJECT_ROOT / "input_data" / "base" / "transformed.csv"
RISK_PATH = PROJECT_ROOT / "input_data" / "risk" / "risk_scores.csv"
REPLACEMENTS_PATH = PROJECT_ROOT / "input_data" / "risk" / "replacement_candidates.csv"
GROUND_TRUTH_MAPPED_PATH = PROJECT_ROOT / "input_data" / "risk" / "ground_truth_replacements.mapped.csv"
GROUND_TRUTH_RAW_PATH = PROJECT_ROOT / "input_data" / "risk" / "ground_truth_replacements.csv"
INVENTORY_SUMMARY_PATH = PROJECT_ROOT / "input_data" / "inventory" / "spare_inventory_cost_summary.csv"
FINANCIAL_COMPARISON_PATH = PROJECT_ROOT / "output_data" / "financial" / "npv_irr_comparison.csv"
FINANCIAL_TIMELINE_PATH = PROJECT_ROOT / "output_data" / "financial" / "monthly_operational_savings.csv"
INVENTORY_CONFIG_PATH = PROJECT_ROOT / "config" / "spare_inventory.json"

OUTPUT_DIR = PROJECT_ROOT / "output_data" / "experimental_results"

STYLE = "seaborn-v0_8-whitegrid"
STL_PERIOD = 7
DEFAULT_LEAD_WINDOW_DAYS = 30
DEFAULT_EPISODE_GAP_DAYS = 7
DEFAULT_PERIODIC_INTERVAL_DAYS = 30
DEFAULT_WATCHLIST_TOP_K = 10
DEFAULT_TUNED_RISK_SCORE = 13.0
DEFAULT_TUNED_RELATIVE_PR = 0.85


@dataclass(frozen=True)
class MethodResult:
    method: str
    alert_episodes: int
    false_alert_episodes: int
    detected_events: int
    total_events: int
    precision: float
    recall: float
    f1: float
    avg_early_days: float | None
    median_early_days: float | None
    alerts_per_detected_event: float | None


def _events_per_100_alerts(detected_events: int, alert_episodes: int) -> float:
    return _safe_ratio(detected_events * 100.0, alert_episodes)


def normalize_device_name(value: object) -> str:
    if value is None or (isinstance(value, float) and math.isnan(value)):
        return ""
    text = str(value).strip()
    text = re.sub(r"\s+", " ", text)
    text = re.sub(r"\bHF0*(\d+)\b", lambda m: f"HF{int(m.group(1))}", text, flags=re.IGNORECASE)
    text = re.sub(
        r"\bInverter\s+0*(\d+)\b",
        lambda m: f"Inverter {int(m.group(1))}",
        text,
        flags=re.IGNORECASE,
    )
    return text


def _safe_ratio(numerator: float, denominator: float) -> float:
    if denominator == 0:
        return 0.0
    return numerator / denominator


def _read_csv(path: Path, **kwargs) -> pd.DataFrame:
    if not path.exists():
        return pd.DataFrame()
    return pd.read_csv(path, **kwargs)


def load_transformed(path: Path = TRANSFORMED_PATH) -> pd.DataFrame:
    transformed = _read_csv(path, parse_dates=["date"])
    if transformed.empty:
        raise FileNotFoundError(f"Missing transformed dataset: {path}")

    rename_map = {}
    if "capacity_kw" in transformed.columns and "installed_capacity_kwp" not in transformed.columns:
        rename_map["capacity_kw"] = "installed_capacity_kwp"
    if "global_irradiation_kwh_m2" in transformed.columns and "irradiation_kwh_m2" not in transformed.columns:
        rename_map["global_irradiation_kwh_m2"] = "irradiation_kwh_m2"
    if rename_map:
        transformed = transformed.rename(columns=rename_map)

    transformed["device_key"] = transformed["device_name"].map(normalize_device_name)
    return transformed


def load_risk(path: Path = RISK_PATH) -> pd.DataFrame:
    risk = _read_csv(path, parse_dates=["risk_date", "latest_date"])
    if risk.empty:
        raise FileNotFoundError(f"Missing risk score dataset: {path}")
    risk["device_key"] = risk["device_name"].map(normalize_device_name)
    return risk


def load_ground_truth(transformed: pd.DataFrame) -> pd.DataFrame:
    path = GROUND_TRUTH_MAPPED_PATH if GROUND_TRUTH_MAPPED_PATH.exists() else GROUND_TRUTH_RAW_PATH
    truth = _read_csv(path)
    if truth.empty:
        return truth

    date_col = None
    for candidate in ("replacement_date", "risk_date", "date"):
        if candidate in truth.columns:
            date_col = candidate
            break
    if date_col is None:
        return pd.DataFrame()

    truth[date_col] = pd.to_datetime(truth[date_col], errors="coerce")
    truth = truth.dropna(subset=[date_col]).copy()
    truth = truth.rename(columns={date_col: "replacement_date"})
    truth["device_key"] = truth["device_name"].map(normalize_device_name)

    zone_lookup = (
        transformed[["zone", "device_name", "device_key"]]
        .dropna(subset=["device_key"])
        .drop_duplicates(["device_key"])
    )
    truth = truth.merge(zone_lookup[["zone", "device_key"]], on="device_key", how="left")
    truth = truth.drop_duplicates(["device_key", "replacement_date"]).sort_values(
        ["replacement_date", "device_key"]
    )
    return truth.reset_index(drop=True)


def resolve_split(dates: pd.Series) -> dict[str, object]:
    unique_dates = pd.Series(pd.to_datetime(dates).dropna().sort_values().unique())
    if unique_dates.empty:
        raise ValueError("No valid dates available to resolve the experiment split.")

    start_date = pd.Timestamp(unique_dates.min())
    end_date = pd.Timestamp(unique_dates.max())
    span_days = int((end_date - start_date).days) + 1
    exact_three_year_span = 365 * 3

    if span_days >= exact_three_year_span:
        train_end = start_date + pd.Timedelta(days=365 * 2 - 1)
        split_label = "calendar split: first 2 years train, next 1 year test"
    else:
        train_days = max(1, round(span_days * 2 / 3))
        train_end = start_date + pd.Timedelta(days=train_days - 1)
        split_label = "ratio split: first 2/3 train, last 1/3 test (used because data < 3 years)"

    test_start = train_end + pd.Timedelta(days=1)
    train_mask = unique_dates <= train_end
    test_mask = unique_dates >= test_start
    train_dates = unique_dates[train_mask]
    test_dates = unique_dates[test_mask]
    if test_dates.empty:
        test_dates = unique_dates.tail(max(1, len(unique_dates) // 3))
        train_dates = unique_dates.iloc[: len(unique_dates) - len(test_dates)]
        test_start = pd.Timestamp(test_dates.min())
        train_end = pd.Timestamp(train_dates.max())
        split_label += "; adjusted to guarantee a non-empty test window"

    return {
        "split_rule": split_label,
        "date_start": start_date.date().isoformat(),
        "date_end": end_date.date().isoformat(),
        "train_start": pd.Timestamp(train_dates.min()).date().isoformat(),
        "train_end": pd.Timestamp(train_dates.max()).date().isoformat(),
        "test_start": pd.Timestamp(test_dates.min()).date().isoformat(),
        "test_end": pd.Timestamp(test_dates.max()).date().isoformat(),
        "train_days": int(len(train_dates)),
        "test_days": int(len(test_dates)),
        "total_days": int(len(unique_dates)),
    }


def resolve_event_overlap_window(
    risk: pd.DataFrame,
    truth: pd.DataFrame,
    lead_window_days: int,
) -> dict[str, object]:
    if truth.empty:
        return {
            "window_start": None,
            "window_end": None,
            "truth_events": 0,
            "note": "No ground-truth replacements overlap the scored risk period.",
        }
    first_event = pd.Timestamp(truth["replacement_date"].min())
    last_event = pd.Timestamp(truth["replacement_date"].max())
    risk_start = pd.Timestamp(risk["risk_date"].min())
    risk_end = pd.Timestamp(risk["risk_date"].max())
    window_start = max(risk_start, first_event - pd.Timedelta(days=lead_window_days))
    window_end = min(risk_end, last_event)
    return {
        "window_start": window_start.date().isoformat(),
        "window_end": window_end.date().isoformat(),
        "truth_events": int(len(truth)),
        "lead_window_days": int(lead_window_days),
        "note": "Supervised event-based evaluation uses all replacement events that overlap the scored risk period.",
    }


def build_setup_summary(
    transformed: pd.DataFrame,
    risk: pd.DataFrame,
    truth: pd.DataFrame,
    split: dict[str, object],
    lead_window_days: int,
    periodic_interval_days: int,
) -> dict[str, object]:
    scored_devices = risk[["zone", "device_key"]].drop_duplicates().shape[0]
    total_devices = transformed[["zone", "device_key"]].drop_duplicates().shape[0]
    overlapped_truth = truth[
        truth["replacement_date"].between(pd.Timestamp(split["date_start"]), pd.Timestamp(split["date_end"]))
    ]
    test_truth = truth[
        truth["replacement_date"].between(pd.Timestamp(split["test_start"]), pd.Timestamp(split["test_end"]))
    ]
    overlap_eval = resolve_event_overlap_window(risk, overlapped_truth, lead_window_days)

    return {
        "dataset": {
            "transformed_rows": int(len(transformed)),
            "risk_rows": int(len(risk)),
            "scored_devices": int(scored_devices),
            "total_devices": int(total_devices),
            "truth_events_total": int(len(truth)),
            "truth_events_in_scored_window": int(len(overlapped_truth)),
            "truth_events_in_test_window": int(len(test_truth)),
        },
        "split": split,
        "supervised_evaluation_window": overlap_eval,
        "metrics": {
            "classification": ["precision", "recall", "f1_score"],
            "lead_time": f"early detection days before confirmed replacement within {lead_window_days}-day match window",
            "ranking": ["recall_at_k", "best_daily_rank", "best_risk_score_before_event"],
            "forecast_fit": ["mae", "rmse"],
            "financial": ["npv", "irr", "dscr"],
        },
        "baselines": {
            "proposed_method": (
                f"tuned supervised operating point: risk_score >= {DEFAULT_TUNED_RISK_SCORE:g} "
                f"and latest_relative_pr < {DEFAULT_TUNED_RELATIVE_PR:.2f}"
            ),
            "proposed_strict_rule": "daily replacement_candidate from composite operational rule",
            "threshold_rule": "alert when latest_relative_pr < 0.70 or latest_pr < 0.40",
            "iforest_rule": "Isolation-Forest-style anomaly score on PR, relative PR, peer deficit, and 30-day trend",
            "periodic_inspection_rule": f"fixed inspection every {periodic_interval_days} days per device",
            "proposed_watchlist_rule": f"daily Top-{DEFAULT_WATCHLIST_TOP_K} devices ranked by risk score",
        },
    }


def write_setup_files(payload: dict[str, object]) -> tuple[Path, Path]:
    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
    json_path = OUTPUT_DIR / "experimental_setup.json"
    md_path = OUTPUT_DIR / "experimental_setup.md"

    json_path.write_text(json.dumps(payload, indent=2), encoding="utf-8")

    split = payload["split"]
    dataset = payload["dataset"]
    metrics = payload["metrics"]
    baselines = payload["baselines"]
    lines = [
        "# Experimental Setup",
        "",
        "## Dataset",
        "",
        f"- Transformed rows: {dataset['transformed_rows']:,}",
        f"- Risk-score rows: {dataset['risk_rows']:,}",
        f"- Devices in transformed data: {dataset['total_devices']:,}",
        f"- Devices with scored risk: {dataset['scored_devices']:,}",
        f"- Ground-truth replacements in file: {dataset['truth_events_total']:,}",
        f"- Ground-truth replacements inside the scored period: {dataset['truth_events_in_scored_window']:,}",
        f"- Ground-truth replacements inside the split test window: {dataset['truth_events_in_test_window']:,}",
        "",
        "## Train/Test Split",
        "",
        f"- Rule: {split['split_rule']}",
        f"- Full date range: {split['date_start']} to {split['date_end']}",
        f"- Train window: {split['train_start']} to {split['train_end']} ({split['train_days']} scored days)",
        f"- Test window: {split['test_start']} to {split['test_end']} ({split['test_days']} scored days)",
        "",
        "## Supervised Event Evaluation Window",
        "",
        f"- Event-evaluation window: {payload['supervised_evaluation_window']['window_start']} to {payload['supervised_evaluation_window']['window_end']}",
        f"- Replacement events evaluated: {payload['supervised_evaluation_window']['truth_events']}",
        f"- Note: {payload['supervised_evaluation_window']['note']}",
        "",
        "## Metrics",
        "",
        f"- Classification: {', '.join(metrics['classification'])}",
        f"- Lead time: {metrics['lead_time']}",
        f"- Ranking: {', '.join(metrics['ranking'])}",
        f"- STL fit: {', '.join(metrics['forecast_fit'])}",
        f"- Financial: {', '.join(metrics['financial'])}",
        "",
        "## Baselines",
        "",
    ]
    for key, value in baselines.items():
        lines.append(f"- {key}: {value}")

    md_path.write_text("\n".join(lines) + "\n", encoding="utf-8")
    return json_path, md_path


def compute_stl_metrics(daily_pr: pd.DataFrame) -> tuple[object, dict[str, float]]:
    observed = daily_pr["mean_pr"].to_numpy(dtype=float)
    stl_result = STL(observed, period=STL_PERIOD, robust=True).fit()
    fitted = stl_result.trend + stl_result.seasonal
    baseline = np.full_like(observed, observed.mean())
    metrics = {
        "baseline_mae": float(np.mean(np.abs(observed - baseline))),
        "baseline_rmse": float(np.sqrt(np.mean((observed - baseline) ** 2))),
        "stl_mae": float(np.mean(np.abs(observed - fitted))),
        "stl_rmse": float(np.sqrt(np.mean((observed - fitted) ** 2))),
    }
    metrics["rmse_improvement_pct"] = _safe_ratio(
        metrics["baseline_rmse"] - metrics["stl_rmse"],
        metrics["baseline_rmse"],
    ) * 100
    return stl_result, metrics


def build_figure1_stl(transformed: pd.DataFrame) -> tuple[Path, dict[str, float]]:
    daily_pr = (
        transformed[
            transformed["performance_ratio"].notna()
            & transformed["irradiation_kwh_m2"].notna()
        ]
        .groupby("date", as_index=False)
        .agg(mean_pr=("performance_ratio", "mean"))
        .sort_values("date")
    )
    if len(daily_pr) < STL_PERIOD * 2:
        raise ValueError("Not enough valid PR points to generate STL decomposition.")

    stl_result, metrics = compute_stl_metrics(daily_pr)

    plt.style.use(STYLE)
    fig, axes = plt.subplots(4, 1, figsize=(13, 10), sharex=True)
    fig.suptitle("Figure 1. STL Decomposition of Daily Mean PR", fontsize=14, fontweight="bold")

    components = [
        ("Observed", stl_result.observed, "#2c7bb6"),
        ("Trend", stl_result.trend, "#d7191c"),
        ("Seasonal", stl_result.seasonal, "#1a9641"),
        ("Residual", stl_result.resid, "#756bb1"),
    ]
    for ax, (label, values, color) in zip(axes, components):
        if label == "Residual":
            ax.bar(daily_pr["date"], values, color=color, alpha=0.7, width=0.9)
            ax.axhline(0, color="#444444", linewidth=0.9, linestyle="--")
        else:
            ax.plot(daily_pr["date"], values, color=color, linewidth=1.5)
        ax.set_ylabel(label)
        ax.xaxis.set_major_formatter(mdates.DateFormatter("%Y-%m"))
        ax.tick_params(axis="x", rotation=30)

    text = (
        f"Naive RMSE: {metrics['baseline_rmse']:.4f}\n"
        f"STL RMSE: {metrics['stl_rmse']:.4f}\n"
        f"RMSE improvement: {metrics['rmse_improvement_pct']:.1f}%"
    )
    axes[0].text(
        0.01,
        0.97,
        text,
        transform=axes[0].transAxes,
        va="top",
        fontsize=8,
        bbox=dict(boxstyle="round,pad=0.4", facecolor="lightyellow", alpha=0.9),
    )

    fig.tight_layout(rect=[0, 0, 1, 0.96])
    out_path = OUTPUT_DIR / "figure1_stl_decomposition.png"
    fig.savefig(out_path, dpi=160, bbox_inches="tight")
    plt.close(fig)
    return out_path, metrics


def collapse_alert_episodes(alerts: pd.DataFrame, gap_days: int) -> pd.DataFrame:
    if alerts.empty:
        return alerts
    alerts = alerts.sort_values(["device_key", "alert_date"]).copy()
    episodes: list[dict[str, object]] = []
    for device_key, group in alerts.groupby("device_key", sort=False):
        last_kept: pd.Timestamp | None = None
        zone = group["zone"].iloc[0] if "zone" in group.columns else None
        for alert_date in pd.to_datetime(group["alert_date"]).sort_values():
            if last_kept is None or (alert_date - last_kept).days > gap_days:
                episodes.append(
                    {
                        "zone": zone,
                        "device_key": device_key,
                        "alert_date": pd.Timestamp(alert_date),
                    }
                )
                last_kept = pd.Timestamp(alert_date)
    return pd.DataFrame(episodes)


def build_threshold_alerts(risk_test: pd.DataFrame) -> pd.DataFrame:
    mask = (risk_test["latest_relative_pr"] < 0.70) | (risk_test["latest_pr"] < 0.40)
    alerts = risk_test.loc[mask, ["zone", "device_key", "risk_date"]].copy()
    return alerts.rename(columns={"risk_date": "alert_date"})


def build_proposed_alerts(risk_test: pd.DataFrame) -> pd.DataFrame:
    alerts = risk_test.loc[risk_test["replacement_candidate"], ["zone", "device_key", "risk_date"]].copy()
    return alerts.rename(columns={"risk_date": "alert_date"})


def build_proposed_supervised_alerts(
    risk_window: pd.DataFrame,
    risk_score_min: float = DEFAULT_TUNED_RISK_SCORE,
    latest_relative_pr_max: float = DEFAULT_TUNED_RELATIVE_PR,
) -> pd.DataFrame:
    mask = (
        (risk_window["risk_score"] >= risk_score_min)
        & (risk_window["latest_relative_pr"] < latest_relative_pr_max)
    )
    alerts = risk_window.loc[mask, ["zone", "device_key", "risk_date"]].copy()
    return alerts.rename(columns={"risk_date": "alert_date"})


def build_topk_watchlist_alerts(risk_window: pd.DataFrame, top_k: int) -> pd.DataFrame:
    rows: list[dict[str, object]] = []
    for risk_date, day in risk_window.groupby("risk_date", sort=True):
        top = day.sort_values(
            ["risk_score", "fleet_relative_risk_score"],
            ascending=[False, False],
        ).head(top_k)
        for _, row in top.iterrows():
            rows.append(
                {
                    "zone": row["zone"],
                    "device_key": row["device_key"],
                    "alert_date": pd.Timestamp(risk_date),
                }
            )
    return pd.DataFrame(rows)


def build_periodic_alerts(risk_test: pd.DataFrame, interval_days: int) -> pd.DataFrame:
    rows: list[dict[str, object]] = []
    ordered = risk_test.sort_values(["zone", "device_key", "risk_date"])
    for (zone, device_key), group in ordered.groupby(["zone", "device_key"], sort=False):
        group = group.sort_values("risk_date")
        dates = pd.to_datetime(group["risk_date"]).drop_duplicates().tolist()
        for idx, alert_date in enumerate(dates):
            if idx % interval_days == 0:
                rows.append({"zone": zone, "device_key": device_key, "alert_date": alert_date})
    return pd.DataFrame(rows)


def build_iforest_alerts(risk_test: pd.DataFrame, contamination: float = 0.05) -> pd.DataFrame:
    feature_cols = [
        "latest_pr",
        "latest_relative_pr",
        "latest_peer_deficit",
        "mean_deficit_14d",
        "severe_low_rate_14d",
        "relative_pr_slope_30d",
    ]
    available_cols = [col for col in feature_cols if col in risk_test.columns]
    if len(available_cols) < 3:
        return pd.DataFrame(columns=["zone", "device_key", "alert_date"])

    work = risk_test[["zone", "device_key", "risk_date"] + available_cols].copy()
    work = work.dropna(subset=available_cols)
    if work.empty:
        return pd.DataFrame(columns=["zone", "device_key", "alert_date"])

    X = work[available_cols].astype(float).to_numpy()
    X = np.nan_to_num(X, nan=0.0)

    try:
        from sklearn.ensemble import IsolationForest
        from sklearn.preprocessing import StandardScaler

        scaler = StandardScaler()
        X_scaled = scaler.fit_transform(X)
        model = IsolationForest(
            contamination=contamination,
            random_state=42,
            n_estimators=200,
        )
        model.fit(X_scaled)
        anomaly_mask = model.predict(X_scaled) == -1
    except Exception:
        # Fallback if scikit-learn is unavailable in the local environment.
        z = (X - X.mean(axis=0)) / np.where(X.std(axis=0) == 0, 1.0, X.std(axis=0))
        score = np.sqrt((z**2).sum(axis=1))
        cutoff = np.quantile(score, 1 - contamination)
        anomaly_mask = score >= cutoff

    alerts = work.loc[anomaly_mask, ["zone", "device_key", "risk_date"]].copy()
    return alerts.rename(columns={"risk_date": "alert_date"})


def evaluate_alerts(
    method: str,
    alerts: pd.DataFrame,
    truth: pd.DataFrame,
    lead_window_days: int,
) -> MethodResult:
    if truth.empty:
        return MethodResult(method, 0, 0, 0, 0, 0.0, 0.0, 0.0, None, None, None)

    truth_sorted = truth.sort_values(["device_key", "replacement_date"]).copy()
    alerts_sorted = alerts.sort_values(["device_key", "alert_date"]).copy()
    alerts_sorted["matched"] = False
    matched_leads: list[int] = []
    matched_events = 0

    for idx, event in truth_sorted.iterrows():
        device_alerts = alerts_sorted[
            (alerts_sorted["device_key"] == event["device_key"])
            & (~alerts_sorted["matched"])
            & (alerts_sorted["alert_date"] <= event["replacement_date"])
        ].copy()
        if device_alerts.empty:
            continue
        device_alerts["lead_days"] = (
            event["replacement_date"] - device_alerts["alert_date"]
        ).dt.days
        candidates = device_alerts[
            (device_alerts["lead_days"] >= 0) & (device_alerts["lead_days"] <= lead_window_days)
        ].sort_values(["lead_days", "alert_date"], ascending=[False, True])
        if candidates.empty:
            continue
        match_idx = candidates.index[0]
        alerts_sorted.loc[match_idx, "matched"] = True
        matched_events += 1
        matched_leads.append(int(candidates.loc[match_idx, "lead_days"]))

    tp = matched_events
    fp = int((~alerts_sorted["matched"]).sum())
    fn = int(len(truth_sorted) - matched_events)
    precision = _safe_ratio(tp, tp + fp)
    recall = _safe_ratio(tp, tp + fn)
    f1 = _safe_ratio(2 * precision * recall, precision + recall) if (precision + recall) else 0.0

    return MethodResult(
        method=method,
        alert_episodes=int(len(alerts_sorted)),
        false_alert_episodes=int(fp),
        detected_events=int(tp),
        total_events=int(len(truth_sorted)),
        precision=float(precision),
        recall=float(recall),
        f1=float(f1),
        avg_early_days=float(np.mean(matched_leads)) if matched_leads else None,
        median_early_days=float(np.median(matched_leads)) if matched_leads else None,
        alerts_per_detected_event=(_safe_ratio(len(alerts_sorted), tp) if tp > 0 else None),
    )


def build_event_ranking_tables(
    risk: pd.DataFrame,
    truth: pd.DataFrame,
    lead_window_days: int,
) -> tuple[Path, Path, Path, pd.DataFrame, pd.DataFrame, pd.DataFrame]:
    diagnostics: list[dict[str, object]] = []
    top_k_levels = [1, 5, 10, 20, 30]

    for _, event in truth.sort_values(["replacement_date", "device_key"]).iterrows():
        event_window = risk[
            (risk["risk_date"] <= event["replacement_date"])
            & (risk["risk_date"] >= event["replacement_date"] - pd.Timedelta(days=lead_window_days))
        ].copy()
        event_ranks: list[tuple[pd.Timestamp, int, float]] = []
        for risk_date, day in event_window.groupby("risk_date", sort=True):
            ranked = day.sort_values(
                ["risk_score", "fleet_relative_risk_score"],
                ascending=[False, False],
            ).reset_index(drop=True)
            match = ranked[ranked["device_key"] == event["device_key"]]
            if match.empty:
                continue
            best_row = match.iloc[0]
            rank = int(match.index[0]) + 1
            event_ranks.append((pd.Timestamp(risk_date), rank, float(best_row["risk_score"])))

        if event_ranks:
            best_rank_day, best_rank, _ = min(event_ranks, key=lambda item: item[1])
            best_score_day, _, best_score = max(event_ranks, key=lambda item: item[2])
            best_rank_lead = int((event["replacement_date"] - best_rank_day).days)
        else:
            best_rank_day = pd.NaT
            best_rank = None
            best_score_day = pd.NaT
            best_score = None
            best_rank_lead = None

        row = {
            "device_name": event["device_key"],
            "replacement_date": pd.Timestamp(event["replacement_date"]).date().isoformat(),
            "best_rank_in_window": best_rank,
            "best_rank_date": None if pd.isna(best_rank_day) else best_rank_day.date().isoformat(),
            "lead_days_at_best_rank": best_rank_lead,
            "best_risk_score_in_window": best_score,
            "best_score_date": None if pd.isna(best_score_day) else best_score_day.date().isoformat(),
        }
        for k in top_k_levels:
            row[f"top_{k}_hit"] = bool(best_rank is not None and best_rank <= k)
        diagnostics.append(row)

    details = pd.DataFrame(diagnostics)
    summary_rows = []
    for k in top_k_levels:
        summary_rows.append(
            {
                "metric": f"recall_at_top_{k}",
                "value": round(float(details[f'top_{k}_hit'].mean()) if not details.empty else 0.0, 4),
                "events_detected": int(details[f"top_{k}_hit"].sum()) if not details.empty else 0,
                "total_events": int(len(details)),
            }
        )

    summary_rows.extend(
        [
            {
                "metric": "median_best_rank",
                "value": round(float(details["best_rank_in_window"].dropna().median()), 2) if details["best_rank_in_window"].notna().any() else None,
                "events_detected": int(details["best_rank_in_window"].notna().sum()),
                "total_events": int(len(details)),
            },
            {
                "metric": "mean_best_rank",
                "value": round(float(details["best_rank_in_window"].dropna().mean()), 2) if details["best_rank_in_window"].notna().any() else None,
                "events_detected": int(details["best_rank_in_window"].notna().sum()),
                "total_events": int(len(details)),
            },
        ]
    )
    summary = pd.DataFrame(summary_rows)
    case_studies = (
        details.sort_values(
            ["top_10_hit", "top_20_hit", "best_rank_in_window", "lead_days_at_best_rank"],
            ascending=[False, False, True, False],
        )
        .head(5)
        .copy()
    )

    summary_path = OUTPUT_DIR / "table1b_event_ranking_summary.csv"
    details_path = OUTPUT_DIR / "event_ranking_diagnostics.csv"
    case_path = OUTPUT_DIR / "table1c_event_case_studies.csv"
    summary.to_csv(summary_path, index=False, encoding="utf-8-sig")
    details.to_csv(details_path, index=False, encoding="utf-8-sig")
    case_studies.to_csv(case_path, index=False, encoding="utf-8-sig")
    return summary_path, details_path, case_path, summary, details, case_studies


def build_table1_and_figure2(
    risk: pd.DataFrame,
    truth: pd.DataFrame,
    eval_window: dict[str, object],
    lead_window_days: int,
    periodic_interval_days: int,
    episode_gap_days: int,
) -> tuple[Path, Path, pd.DataFrame]:
    test_start = pd.Timestamp(eval_window["window_start"])
    test_end = pd.Timestamp(eval_window["window_end"])

    risk_test = risk[risk["risk_date"].between(test_start, test_end)].copy()
    truth_test = truth[truth["replacement_date"].between(test_start, test_end)].copy()

    method_alert_builders = [
        ("Proposed AI", build_proposed_supervised_alerts(risk_test)),
        ("Proposed AI (strict operational rule)", build_proposed_alerts(risk_test)),
        (f"Proposed Top-{DEFAULT_WATCHLIST_TOP_K} Watchlist", build_topk_watchlist_alerts(risk_test, DEFAULT_WATCHLIST_TOP_K)),
        ("Threshold baseline", build_threshold_alerts(risk_test)),
        ("Isolation Forest", build_iforest_alerts(risk_test)),
        ("Periodic inspection", build_periodic_alerts(risk_test, periodic_interval_days)),
    ]

    results: list[MethodResult] = []
    for method_name, alerts in method_alert_builders:
        collapsed = collapse_alert_episodes(alerts, gap_days=episode_gap_days)
        results.append(
            evaluate_alerts(
                method=method_name,
                alerts=collapsed,
                truth=truth_test,
                lead_window_days=lead_window_days,
            )
        )

    table = pd.DataFrame(
        [
            {
                "method": result.method,
                "alert_episodes": result.alert_episodes,
                "false_alert_episodes": result.false_alert_episodes,
                "alerts_per_detected_event": None if result.alerts_per_detected_event is None else round(result.alerts_per_detected_event, 2),
                "detected_events": result.detected_events,
                "total_events": result.total_events,
                "precision": round(result.precision, 4),
                "recall": round(result.recall, 4),
                "f1_score": round(result.f1, 4),
                "avg_early_days": None if result.avg_early_days is None else round(result.avg_early_days, 2),
                "median_early_days": None if result.median_early_days is None else round(result.median_early_days, 2),
            }
            for result in results
        ]
    )
    table = table.sort_values(["f1_score", "recall", "precision"], ascending=False).reset_index(drop=True)

    table_path = OUTPUT_DIR / "table1_anomaly_baseline_comparison.csv"
    table.to_csv(table_path, index=False, encoding="utf-8-sig")

    plt.style.use(STYLE)
    fig, axes = plt.subplots(1, 3, figsize=(16, 5.5))
    x = np.arange(len(table))
    width = 0.24

    axes[0].bar(x - width, table["precision"], width=width, label="Precision", color="#2c7fb8")
    axes[0].bar(x, table["recall"], width=width, label="Recall", color="#41ab5d")
    axes[0].bar(x + width, table["f1_score"], width=width, label="F1", color="#d7301f")
    axes[0].set_xticks(x)
    axes[0].set_xticklabels(table["method"], rotation=20, ha="right")
    axes[0].set_ylim(0, max(1.0, float(table[["precision", "recall", "f1_score"]].max().max()) * 1.15))
    axes[0].set_ylabel("Score")
    axes[0].set_title("Detection Metrics")
    axes[0].legend(fontsize=8)

    lead_values = table["avg_early_days"].fillna(0)
    axes[1].bar(x, lead_values, color="#756bb1", alpha=0.85)
    axes[1].set_xticks(x)
    axes[1].set_xticklabels(table["method"], rotation=20, ha="right")
    axes[1].set_ylabel("Days earlier than replacement")
    axes[1].set_title("Average Early Warning")

    burden_values = table["alert_episodes"].fillna(0)
    axes[2].bar(x, burden_values, color="#636363", alpha=0.85)
    axes[2].set_xticks(x)
    axes[2].set_xticklabels(table["method"], rotation=20, ha="right")
    axes[2].set_ylabel("Alert episodes")
    axes[2].set_yscale("log")
    axes[2].set_title("Operational Alert Burden (log scale)")

    fig.suptitle("Figure 2. Event-Based Baseline Comparison on Overlapped Replacement Events", fontsize=14, fontweight="bold")
    fig.tight_layout(rect=[0, 0, 1, 0.95])
    figure_path = OUTPUT_DIR / "figure2_baseline_comparison.png"
    fig.savefig(figure_path, dpi=160, bbox_inches="tight")
    plt.close(fig)
    return table_path, figure_path, table


def build_primary_metrics_table(table1: pd.DataFrame) -> tuple[Path, pd.DataFrame]:
    primary = table1.copy()
    primary["event_coverage"] = primary["detected_events"] / primary["total_events"]
    primary["events_per_100_alerts"] = primary.apply(
        lambda row: round(_events_per_100_alerts(int(row["detected_events"]), int(row["alert_episodes"])), 2),
        axis=1,
    )
    primary = primary[
        [
            "method",
            "detected_events",
            "total_events",
            "event_coverage",
            "median_early_days",
            "alert_episodes",
            "events_per_100_alerts",
            "alerts_per_detected_event",
        ]
    ].copy()
    primary["event_coverage"] = primary["event_coverage"].round(4)
    path = OUTPUT_DIR / "table1_primary_metrics.csv"
    primary.to_csv(path, index=False, encoding="utf-8-sig")
    return path, primary


def build_table2_inventory() -> tuple[Path, pd.DataFrame]:
    curve = _read_csv(INVENTORY_SUMMARY_PATH)
    if curve.empty:
        raise FileNotFoundError(f"Missing inventory summary: {INVENTORY_SUMMARY_PATH}")

    reactive = curve.loc[curve["stock_n"] == 0].iloc[0]
    prestock = curve.loc[curve["total_cost_vnd"].idxmin()]

    rows = [
        {
            "scenario": "Reactive",
            "stock_n": int(reactive["stock_n"]),
            "expected_generation_loss_vnd": float(reactive["expected_generation_loss_vnd"]),
            "holding_cost_vnd": float(reactive["holding_cost_vnd"]),
            "total_cost_vnd": float(reactive["total_cost_vnd"]),
            "saving_vs_reactive_vnd": 0.0,
            "cost_reduction_pct_vs_reactive": 0.0,
        },
        {
            "scenario": "Pre-stock optimal",
            "stock_n": int(prestock["stock_n"]),
            "expected_generation_loss_vnd": float(prestock["expected_generation_loss_vnd"]),
            "holding_cost_vnd": float(prestock["holding_cost_vnd"]),
            "total_cost_vnd": float(prestock["total_cost_vnd"]),
            "saving_vs_reactive_vnd": float(reactive["total_cost_vnd"] - prestock["total_cost_vnd"]),
            "cost_reduction_pct_vs_reactive": _safe_ratio(
                float(reactive["total_cost_vnd"] - prestock["total_cost_vnd"]),
                float(reactive["total_cost_vnd"]),
            )
            * 100,
        },
    ]
    table = pd.DataFrame(rows)
    table_path = OUTPUT_DIR / "table2_inventory_cost_benefit.csv"
    table.to_csv(table_path, index=False, encoding="utf-8-sig")
    return table_path, table


def try_load_workbook_assumptions() -> dict[str, float]:
    workbook_path = PROJECT_ROOT / "input_data" / "Financial_Assumptions.xlsx"
    if not workbook_path.exists():
        return {}

    try:
        from openpyxl import load_workbook
    except Exception:
        return {}

    wb = load_workbook(workbook_path, data_only=True)
    if "01_Assumptions" not in wb.sheetnames:
        return {}
    ws = wb["01_Assumptions"]

    assumptions: dict[str, float] = {}
    for row in ws.iter_rows(min_row=4, values_only=True):
        if not row or not row[0]:
            continue
        key = str(row[0]).strip()
        raw = row[1]
        if raw in (None, ""):
            continue
        if isinstance(raw, str):
            text = raw.strip().replace(",", "")
            is_pct = text.endswith("%")
            if is_pct:
                text = text[:-1]
            try:
                value = float(text)
            except ValueError:
                continue
            assumptions[key] = value / 100 if is_pct else value
        else:
            try:
                assumptions[key] = float(raw)
            except (TypeError, ValueError):
                continue
    return assumptions


def compute_dscr_rows(financial_timeline: pd.DataFrame) -> tuple[float | None, float | None]:
    assumptions = try_load_workbook_assumptions()
    debt_service = None
    for key in (
        "annual_debt_service_vnd",
        "debt_service_annual_vnd",
        "annual_debt_service",
        "debt_service_annual",
    ):
        if key in assumptions:
            debt_service = assumptions[key]
            break

    if debt_service is None or debt_service == 0 or financial_timeline.empty:
        return None, None

    annual_delta = float(financial_timeline["net_cost_delta_excl_revenue"].sum())
    baseline_cfads = assumptions.get("cfads_without_ai_vnd") or assumptions.get("cfads_baseline_vnd")
    if baseline_cfads is None:
        return None, None

    before = baseline_cfads / debt_service
    after = (baseline_cfads + annual_delta) / debt_service
    return before, after


def quarterly_to_annual(irr_quarterly: float) -> float:
    return (1 + irr_quarterly) ** 4 - 1


def build_table3_financial() -> tuple[Path, pd.DataFrame]:
    comparison = _read_csv(FINANCIAL_COMPARISON_PATH)
    if comparison.empty:
        raise FileNotFoundError(f"Missing financial comparison file: {FINANCIAL_COMPARISON_PATH}")
    timeline = _read_csv(FINANCIAL_TIMELINE_PATH)

    before = comparison.loc[comparison["scenario"] == "Without AI"].iloc[0]
    after = comparison.loc[comparison["scenario"] == "With AI"].iloc[0]

    npv_before = float(before["npv_vnd"]) if pd.notna(before["npv_vnd"]) else math.nan
    npv_after = float(after["npv_vnd"]) if pd.notna(after["npv_vnd"]) else math.nan

    if pd.notna(before.get("irr_annual_effective")):
        irr_before = float(before["irr_annual_effective"])
    elif pd.notna(before.get("irr_quarterly")):
        irr_before = quarterly_to_annual(float(before["irr_quarterly"]))
    else:
        irr_before = math.nan

    irr_after = float(after["irr_annual_effective"]) if pd.notna(after.get("irr_annual_effective")) else math.nan
    dscr_before, dscr_after = compute_dscr_rows(timeline)

    rows = [
        {
            "metric": "NPV",
            "before": npv_before,
            "after": npv_after,
            "delta": npv_after - npv_before if not math.isnan(npv_before) and not math.isnan(npv_after) else math.nan,
            "delta_pct": _safe_ratio(npv_after - npv_before, abs(npv_before)) * 100
            if not math.isnan(npv_before) and npv_before != 0 and not math.isnan(npv_after)
            else math.nan,
        },
        {
            "metric": "IRR_annual_effective",
            "before": irr_before,
            "after": irr_after,
            "delta": irr_after - irr_before if not math.isnan(irr_before) and not math.isnan(irr_after) else math.nan,
            "delta_pct": _safe_ratio(irr_after - irr_before, abs(irr_before)) * 100
            if not math.isnan(irr_before) and irr_before != 0 and not math.isnan(irr_after)
            else math.nan,
        },
        {
            "metric": "DSCR",
            "before": dscr_before,
            "after": dscr_after,
            "delta": (dscr_after - dscr_before) if dscr_before is not None and dscr_after is not None else math.nan,
            "delta_pct": _safe_ratio(dscr_after - dscr_before, abs(dscr_before)) * 100
            if dscr_before not in (None, 0) and dscr_after is not None
            else math.nan,
        },
    ]
    table = pd.DataFrame(rows)
    table_path = OUTPUT_DIR / "table3_financial_impact.csv"
    table.to_csv(table_path, index=False, encoding="utf-8-sig")
    return table_path, table


def write_summary_report(
    setup: dict[str, object],
    stl_metrics: dict[str, float],
    table1: pd.DataFrame,
    primary_metrics: pd.DataFrame,
    ranking_summary: pd.DataFrame,
    table2: pd.DataFrame,
    table3: pd.DataFrame,
) -> Path:
    best_method = table1.sort_values(["f1_score", "precision", "recall"], ascending=False).iloc[0]
    watchlist_row = table1[table1["method"] == f"Proposed Top-{DEFAULT_WATCHLIST_TOP_K} Watchlist"]
    inventory_best = table2.loc[table2["scenario"] == "Pre-stock optimal"].iloc[0]
    recall_at_10 = ranking_summary.loc[ranking_summary["metric"] == "recall_at_top_10", "value"]
    median_best_rank = ranking_summary.loc[ranking_summary["metric"] == "median_best_rank", "value"]
    primary_best = primary_metrics[primary_metrics["method"] == "Proposed AI"]
    lines = [
        "# Experimental Results Summary",
        "",
        "## 5.1 Experimental Setup",
        "",
        f"- Split rule: {setup['split']['split_rule']}",
        f"- Train window: {setup['split']['train_start']} to {setup['split']['train_end']}",
        f"- Test window: {setup['split']['test_start']} to {setup['split']['test_end']}",
        f"- Ground-truth events in scored period: {setup['dataset']['truth_events_in_scored_window']}",
        f"- Supervised event-evaluation window: {setup['supervised_evaluation_window']['window_start']} to {setup['supervised_evaluation_window']['window_end']}",
        "",
        "## 5.2 STL Decomposition",
        "",
        f"- STL RMSE: {stl_metrics['stl_rmse']:.4f}",
        f"- Naive RMSE: {stl_metrics['baseline_rmse']:.4f}",
        f"- RMSE improvement: {stl_metrics['rmse_improvement_pct']:.2f}%",
        "",
        "## 5.3-5.4 Anomaly Detection and Baseline Comparison",
        "",
        f"- Best method by F1: {best_method['method']}",
        f"- Precision / Recall / F1: {best_method['precision']:.4f} / {best_method['recall']:.4f} / {best_method['f1_score']:.4f}",
        f"- Average early warning: {best_method['avg_early_days'] if pd.notna(best_method['avg_early_days']) else 'N/A'} days",
        f"- Proposed AI events per 100 alerts: {primary_best.iloc[0]['events_per_100_alerts']:.2f}" if not primary_best.empty else "- Proposed AI events per 100 alerts: N/A",
        f"- Proposed AI median lead time: {primary_best.iloc[0]['median_early_days']:.2f} days" if not primary_best.empty and pd.notna(primary_best.iloc[0]['median_early_days']) else "- Proposed AI median lead time: N/A",
        f"- Proposed Top-{DEFAULT_WATCHLIST_TOP_K} watchlist Recall / F1: {watchlist_row.iloc[0]['recall']:.4f} / {watchlist_row.iloc[0]['f1_score']:.4f}" if not watchlist_row.empty else "- Proposed Top-watchlist metrics: N/A",
        f"- Proposed ranking Recall@10: {float(recall_at_10.iloc[0]):.4f}" if not recall_at_10.empty else "- Proposed ranking Recall@10: N/A",
        f"- Proposed median best rank before replacement: {float(median_best_rank.iloc[0]):.2f}" if not median_best_rank.empty and pd.notna(median_best_rank.iloc[0]) else "- Proposed median best rank before replacement: N/A",
        "",
        "## 5.5 Cost-Benefit",
        "",
        f"- Optimal pre-stock level: {int(inventory_best['stock_n'])} spare units",
        f"- Total saving vs reactive: {inventory_best['saving_vs_reactive_vnd']:,.0f} VND",
        f"- Cost reduction vs reactive: {inventory_best['cost_reduction_pct_vs_reactive']:.2f}%",
        "",
        "## 5.6 Financial Impact",
        "",
    ]
    for _, row in table3.iterrows():
        before = row["before"]
        after = row["after"]
        delta_pct = row["delta_pct"]
        before_text = "N/A" if pd.isna(before) else f"{before:,.4f}"
        after_text = "N/A" if pd.isna(after) else f"{after:,.4f}"
        delta_text = "N/A" if pd.isna(delta_pct) else f"{delta_pct:.2f}%"
        lines.append(f"- {row['metric']}: {before_text} -> {after_text} (delta {delta_text})")

    lines.extend(
        [
            "",
            "## Artifact Paths",
            "",
            "- `output_data/experimental_results/figure1_stl_decomposition.png`",
            "- `output_data/experimental_results/figure2_baseline_comparison.png`",
            "- `output_data/experimental_results/table1_anomaly_baseline_comparison.csv`",
            "- `output_data/experimental_results/table1b_event_ranking_summary.csv`",
            "- `output_data/experimental_results/table1_primary_metrics.csv`",
            "- `output_data/experimental_results/table1c_event_case_studies.csv`",
            "- `output_data/experimental_results/event_ranking_diagnostics.csv`",
            "- `output_data/experimental_results/table2_inventory_cost_benefit.csv`",
            "- `output_data/experimental_results/table3_financial_impact.csv`",
        ]
    )
    path = OUTPUT_DIR / "results_summary.md"
    path.write_text("\n".join(lines) + "\n", encoding="utf-8")
    return path


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(description="Generate paper-ready experimental result tables and figures.")
    parser.add_argument("--lead-window-days", type=int, default=DEFAULT_LEAD_WINDOW_DAYS)
    parser.add_argument("--episode-gap-days", type=int, default=DEFAULT_EPISODE_GAP_DAYS)
    parser.add_argument("--periodic-interval-days", type=int, default=DEFAULT_PERIODIC_INTERVAL_DAYS)
    return parser.parse_args()


def main() -> None:
    args = parse_args()
    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)

    transformed = load_transformed()
    risk = load_risk()
    truth = load_ground_truth(transformed)
    split = resolve_split(risk["risk_date"])
    setup = build_setup_summary(
        transformed=transformed,
        risk=risk,
        truth=truth,
        split=split,
        lead_window_days=args.lead_window_days,
        periodic_interval_days=args.periodic_interval_days,
    )

    setup_json, setup_md = write_setup_files(setup)
    figure1_path, stl_metrics = build_figure1_stl(transformed)
    eval_truth = truth[
        truth["replacement_date"].between(
            pd.Timestamp(setup["supervised_evaluation_window"]["window_start"]),
            pd.Timestamp(setup["supervised_evaluation_window"]["window_end"]),
        )
    ].copy()
    ranking_summary_path, ranking_details_path, case_studies_path, ranking_summary, ranking_details, case_studies = build_event_ranking_tables(
        risk=risk,
        truth=eval_truth,
        lead_window_days=args.lead_window_days,
    )
    table1_path, figure2_path, table1 = build_table1_and_figure2(
        risk=risk,
        truth=truth,
        eval_window=setup["supervised_evaluation_window"],
        lead_window_days=args.lead_window_days,
        periodic_interval_days=args.periodic_interval_days,
        episode_gap_days=args.episode_gap_days,
    )
    primary_table_path, primary_metrics = build_primary_metrics_table(table1)
    table2_path, table2 = build_table2_inventory()
    table3_path, table3 = build_table3_financial()
    summary_path = write_summary_report(setup, stl_metrics, table1, primary_metrics, ranking_summary, table2, table3)

    print(f"Experimental setup JSON -> {setup_json}")
    print(f"Experimental setup Markdown -> {setup_md}")
    print(f"Figure 1 -> {figure1_path}")
    print(f"Figure 2 -> {figure2_path}")
    print(f"Table 1 -> {table1_path}")
    print(f"Table 1b -> {ranking_summary_path}")
    print(f"Table 1 primary -> {primary_table_path}")
    print(f"Table 1c -> {case_studies_path}")
    print(f"Event diagnostics -> {ranking_details_path}")
    print(f"Table 2 -> {table2_path}")
    print(f"Table 3 -> {table3_path}")
    print(f"Summary -> {summary_path}")


if __name__ == "__main__":
    main()
